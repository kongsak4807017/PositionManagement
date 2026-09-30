from __future__ import annotations

import argparse
import hashlib
import os
import re
from datetime import date, datetime, timedelta, timezone
from enum import Enum as PyEnum
from pathlib import Path
from uuid import uuid4

import jwt
from openpyxl import load_workbook
from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm
from jwt import InvalidTokenError
from passlib.context import CryptContext
from pydantic import BaseModel, Field
from sqlalchemy import (
    JSON, Boolean, Column, Date, DateTime, Enum, ForeignKey, Integer,
    String, Text, UniqueConstraint, create_engine, func, or_,
)
from sqlalchemy.orm import Session, declarative_base, sessionmaker


APP_ENV = os.getenv("APP_ENV", "development").lower()
DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./chro_production.db")
JWT_SECRET = os.getenv("JWT_SECRET", "dev-only-change-me-before-production")
JWT_ALGORITHM = os.getenv("JWT_ALGORITHM", "HS256")
ACCESS_TOKEN_MINUTES = int(os.getenv("ACCESS_TOKEN_MINUTES", "480"))
HROPS_STORAGE_DIR = Path(os.getenv("HROPS_STORAGE_DIR", "./data/hrops"))
MAX_HROPS_UPLOAD_MB = int(os.getenv("MAX_HROPS_UPLOAD_MB", "100"))
ALLOWED_ORIGINS = [
    x.strip() for x in os.getenv(
        "ALLOWED_ORIGINS",
        "http://localhost:8000,http://127.0.0.1:8000",
    ).split(",") if x.strip()
]

if APP_ENV == "production" and (JWT_SECRET == "dev-only-change-me-before-production" or len(JWT_SECRET) < 32):
    raise RuntimeError("JWT_SECRET must be a random value of at least 32 characters in production.")

connect_args = {"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {}
engine = create_engine(DATABASE_URL, pool_pre_ping=True, future=True, connect_args=connect_args)
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)
Base = declarative_base()
pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="/api/auth/login")


def uid() -> str:
    return str(uuid4())


class UserRole(str, PyEnum):
    MOPH_ADMIN = "MOPH_ADMIN"
    REGION_ADMIN = "REGION_ADMIN"
    REGION_EXECUTIVE = "REGION_EXECUTIVE"
    PROVINCE_ADMIN = "PROVINCE_ADMIN"
    HOSPITAL_HR = "HOSPITAL_HR"
    AUDITOR = "AUDITOR"


class VacancyStatus(str, PyEnum):
    OPEN = "OPEN"
    WAITING_APPROVAL = "WAITING_APPROVAL"
    READY_TO_RECRUIT = "READY_TO_RECRUIT"
    RECRUITING = "RECRUITING"
    APPOINTING = "APPOINTING"
    FILLED = "FILLED"
    CANCELLED = "CANCELLED"


class Province(Base):
    __tablename__ = "provinces"
    province_code = Column(String(2), primary_key=True)
    province_name_th = Column(String(100), nullable=False)
    health_region = Column(String(4), nullable=False, default="1")


class OrganizationalUnit(Base):
    __tablename__ = "organizational_units"
    unit_id = Column(String(40), primary_key=True)
    unit_name = Column(String(255), nullable=False)
    unit_type_label = Column(String(30), nullable=False)
    province_code = Column(String(2), ForeignKey("provinces.province_code"), nullable=False, index=True)
    amphur_name = Column(String(120))


class User(Base):
    __tablename__ = "users"
    user_id = Column(String(36), primary_key=True, default=uid)
    username = Column(String(120), nullable=False, unique=True, index=True)
    password_hash = Column(String(255), nullable=False)
    full_name = Column(String(255), nullable=False)
    role = Column(Enum(UserRole, native_enum=False), nullable=False, index=True)
    province_code = Column(String(2), ForeignKey("provinces.province_code"), index=True)
    unit_id = Column(String(40), ForeignKey("organizational_units.unit_id"), index=True)
    is_active = Column(Boolean, nullable=False, default=True)
    created_at = Column(DateTime(timezone=True), nullable=False, server_default=func.now())
    last_login_at = Column(DateTime(timezone=True))


class Position(Base):
    __tablename__ = "positions"
    position_uid = Column(String(36), primary_key=True, default=uid)
    chro_position_id = Column(String(60), nullable=False, unique=True, index=True)
    hrops_position_no = Column(String(80), index=True)
    unit_id = Column(String(40), ForeignKey("organizational_units.unit_id"), nullable=False, index=True)
    position_type = Column(String(120))
    position_name_th = Column(String(255), nullable=False)
    position_level = Column(String(120))
    employment_type = Column(String(120))
    is_active = Column(Boolean, nullable=False, default=True)
    created_at = Column(DateTime(timezone=True), nullable=False, server_default=func.now())
    updated_at = Column(DateTime(timezone=True), nullable=False, server_default=func.now(), onupdate=func.now())


class HropsImportRun(Base):
    __tablename__ = "hrops_import_runs"
    __table_args__ = (UniqueConstraint("baseline_month", "sha256", name="uq_hrops_month_sha"),)
    import_id = Column(String(36), primary_key=True, default=uid)
    baseline_month = Column(Date, nullable=False, index=True)
    original_filename = Column(String(255), nullable=False)
    stored_path = Column(String(500), nullable=False)
    sha256 = Column(String(64), nullable=False, index=True)
    file_size_bytes = Column(Integer, nullable=False)
    status = Column(String(30), nullable=False, default="RECEIVED", index=True)
    row_count = Column(Integer)
    validation_summary = Column(JSON)
    notes = Column(Text)
    uploaded_by_user_id = Column(String(36), ForeignKey("users.user_id"), nullable=False)
    created_at = Column(DateTime(timezone=True), nullable=False, server_default=func.now())


class HropsStagingRow(Base):
    __tablename__ = "hrops_staging_rows"
    __table_args__ = (UniqueConstraint("import_id", "row_no", name="uq_hrops_import_row"),)
    staging_id = Column(String(36), primary_key=True, default=uid)
    import_id = Column(String(36), ForeignKey("hrops_import_runs.import_id"), nullable=False, index=True)
    row_no = Column(Integer, nullable=False)
    raw_payload = Column(JSON, nullable=False)
    validation_status = Column(String(20), nullable=False, default="PENDING")
    validation_errors = Column(JSON)


class HropsPositionSnapshot(Base):
    __tablename__ = "hrops_position_snapshots"
    snapshot_id = Column(String(36), primary_key=True, default=uid)
    import_id = Column(String(36), ForeignKey("hrops_import_runs.import_id"), nullable=False, index=True)
    snapshot_month = Column(Date, nullable=False, index=True)
    position_uid = Column(String(36), ForeignKey("positions.position_uid"), index=True)
    hrops_position_no = Column(String(80), nullable=False, index=True)
    unit_id = Column(String(40), index=True)
    position_type = Column(String(120))
    position_name_th = Column(String(255))
    position_level = Column(String(120))
    holder_status = Column(String(40))
    raw_payload = Column(JSON, nullable=False)
    created_at = Column(DateTime(timezone=True), nullable=False, server_default=func.now())


class VacancyCase(Base):
    __tablename__ = "vacancy_cases"
    case_id = Column(String(36), primary_key=True, default=uid)
    case_no = Column(String(80), nullable=False, unique=True, index=True)
    position_uid = Column(String(36), ForeignKey("positions.position_uid"), nullable=False, index=True)
    responsible_unit_id = Column(String(40), ForeignKey("organizational_units.unit_id"), nullable=False, index=True)
    vacant_date = Column(Date, nullable=False, index=True)
    vacant_reason = Column(String(150), nullable=False)
    status = Column(Enum(VacancyStatus, native_enum=False), nullable=False, default=VacancyStatus.OPEN, index=True)
    current_milestone = Column(String(10), nullable=False, default="M1", index=True)
    retirement_use_approved = Column(Boolean, nullable=False, default=False)
    retirement_approval_doc_no = Column(String(180))
    retirement_approval_doc_date = Column(Date)
    retirement_use_from_date = Column(Date)
    remarks = Column(Text)
    created_by_user_id = Column(String(36), ForeignKey("users.user_id"), nullable=False)
    created_at = Column(DateTime(timezone=True), nullable=False, server_default=func.now())
    updated_at = Column(DateTime(timezone=True), nullable=False, server_default=func.now(), onupdate=func.now())
    closed_at = Column(DateTime(timezone=True))


class VacancyEvent(Base):
    __tablename__ = "vacancy_events"
    event_id = Column(String(36), primary_key=True, default=uid)
    case_id = Column(String(36), ForeignKey("vacancy_cases.case_id"), nullable=False, index=True)
    event_type = Column(String(80), nullable=False, index=True)
    from_status = Column(String(40))
    to_status = Column(String(40))
    milestone = Column(String(10))
    reference_doc_no = Column(String(180))
    notes = Column(Text)
    actor_user_id = Column(String(36), ForeignKey("users.user_id"), nullable=False)
    created_at = Column(DateTime(timezone=True), nullable=False, server_default=func.now(), index=True)


class AuditLog(Base):
    __tablename__ = "audit_logs"
    audit_id = Column(String(36), primary_key=True, default=uid)
    actor_user_id = Column(String(36), ForeignKey("users.user_id"), index=True)
    action = Column(String(120), nullable=False, index=True)
    entity_type = Column(String(80), nullable=False, index=True)
    entity_id = Column(String(80), index=True)
    request_path = Column(String(500))
    ip_address = Column(String(80))
    before_json = Column(JSON)
    after_json = Column(JSON)
    created_at = Column(DateTime(timezone=True), nullable=False, server_default=func.now(), index=True)


class TokenResponse(BaseModel):
    access_token: str
    token_type: str = "bearer"


class UserCreate(BaseModel):
    username: str = Field(min_length=3, max_length=120)
    password: str = Field(min_length=12, max_length=200)
    full_name: str = Field(min_length=2, max_length=255)
    role: UserRole
    province_code: str | None = None
    unit_id: str | None = None


class PositionCreate(BaseModel):
    chro_position_id: str = Field(min_length=3, max_length=60)
    hrops_position_no: str | None = Field(default=None, max_length=80)
    unit_id: str = Field(min_length=2, max_length=40)
    position_type: str | None = Field(default=None, max_length=120)
    position_name_th: str = Field(min_length=2, max_length=255)
    position_level: str | None = Field(default=None, max_length=120)
    employment_type: str | None = Field(default=None, max_length=120)


class VacancyCreate(BaseModel):
    case_no: str = Field(min_length=3, max_length=80)
    position_uid: str
    responsible_unit_id: str
    vacant_date: date
    vacant_reason: str = Field(min_length=2, max_length=150)
    status: VacancyStatus = VacancyStatus.OPEN
    current_milestone: str = Field(default="M1", max_length=10)
    retirement_use_approved: bool = False
    retirement_approval_doc_no: str | None = None
    retirement_approval_doc_date: date | None = None
    retirement_use_from_date: date | None = None
    remarks: str | None = None


class VacancyEventCreate(BaseModel):
    event_type: str = Field(min_length=2, max_length=80)
    to_status: VacancyStatus | None = None
    milestone: str | None = Field(default=None, max_length=10)
    reference_doc_no: str | None = None
    notes: str | None = None
    retirement_use_approved: bool | None = None
    retirement_approval_doc_no: str | None = None
    retirement_approval_doc_date: date | None = None
    retirement_use_from_date: date | None = None


READ_ALL = {UserRole.MOPH_ADMIN, UserRole.REGION_ADMIN, UserRole.REGION_EXECUTIVE, UserRole.AUDITOR}
WRITE_ALL = {UserRole.MOPH_ADMIN, UserRole.REGION_ADMIN}
WRITE = WRITE_ALL | {UserRole.PROVINCE_ADMIN, UserRole.HOSPITAL_HR}


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def hash_password(password: str) -> str:
    if len(password) < 12:
        raise ValueError("Password must be at least 12 characters.")
    return pwd_context.hash(password)


def create_token(user: User) -> str:
    now = datetime.now(timezone.utc)
    return jwt.encode(
        {"sub": user.user_id, "role": user.role.value, "iat": now, "exp": now + timedelta(minutes=ACCESS_TOKEN_MINUTES)},
        JWT_SECRET,
        algorithm=JWT_ALGORITHM,
    )


def current_user(token: str = Depends(oauth2_scheme), db: Session = Depends(get_db)) -> User:
    error = HTTPException(status_code=401, detail="Invalid or expired access token.", headers={"WWW-Authenticate": "Bearer"})
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        user_id = payload.get("sub")
        if not user_id:
            raise error
    except InvalidTokenError:
        raise error
    user = db.get(User, user_id)
    if not user or not user.is_active:
        raise error
    return user


def require_roles(*roles: UserRole):
    allowed = set(roles)
    def dep(user: User = Depends(current_user)) -> User:
        if user.role not in allowed:
            raise HTTPException(status_code=403, detail="Insufficient role permission.")
        return user
    return dep


def assert_unit_scope(db: Session, user: User, unit_id: str, write: bool = False) -> None:
    if write and user.role not in WRITE:
        raise HTTPException(status_code=403, detail="Read-only role.")
    if user.role in READ_ALL:
        return
    unit = db.get(OrganizationalUnit, unit_id)
    if not unit:
        raise HTTPException(status_code=404, detail="Organizational unit not found.")
    if user.role == UserRole.PROVINCE_ADMIN and user.province_code == unit.province_code:
        return
    if user.role == UserRole.HOSPITAL_HR and user.unit_id == unit.unit_id:
        return
    raise HTTPException(status_code=403, detail="Outside data scope.")


def write_audit(db: Session, user: User | None, action: str, entity_type: str, entity_id: str | None, request: Request | None, before=None, after=None):
    db.add(AuditLog(
        actor_user_id=user.user_id if user else None,
        action=action,
        entity_type=entity_type,
        entity_id=entity_id,
        request_path=request.url.path if request else None,
        ip_address=request.client.host if request and request.client else None,
        before_json=before,
        after_json=after,
    ))


J18_DATA_SHEET_INDEX = 1          # Sheet 2: "เขต 1"
J18_TECHNICAL_HEADER_ROW = 6     # Row 5 = Thai labels, Row 6 = stable technical field names

# Mapping verified against "จ18 1 กย 69", Sheet 2 ("เขต 1").
# Position master uses the "ตาม อต." / pos_* organizational fields because vacant
# positions still retain those fields even when person/pay_* fields are blank.
J18_FIELD_MAP = {
    "external_position_id": "position_id",
    "hrops_position_no": "position_code",
    "unit_id": "pos_ou_path3_code",
    "unit_name": "pos_rev_name",
    "unit_type": "pos_ประเภทส่วนราชการ",
    "province_code": "pos_province_code",
    "province_name": "pos_province_name",
    "amphur_name": "pos_amphur_name",
    "position_name_th": "line_position_name",
    "position_type": "position_type_name",
    "position_level_from": "from_level_name",
    "position_level_to": "to_level_name",
    "current_position_level": "pay_position_level_name",
    "employment_type": "ประเภทบุคลากร",
    "holder_status": "สถานะตำแหน่ง",
    "employment_status": "employment_rev_status",
    "retirement_date": "retirement_date",
    "vacancy_date": "empty_date",
    "vacancy_reason": "empty_movement_name",
}

J18_REQUIRED_TECHNICAL_HEADERS = {
    "position_id",
    "position_code",
    "pos_ou_path3_code",
    "pos_rev_name",
    "pos_province_code",
    "line_position_name",
    "position_type_name",
    "from_level_name",
    "to_level_name",
    "สถานะตำแหน่ง",
}


def json_value(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def text_value(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def compose_position_level(record: dict) -> str | None:
    """Use actual current level when occupied; otherwise use the authorized level range."""
    holder_status = text_value(record.get("holder_status")) or ""
    current = text_value(record.get("current_position_level"))
    level_from = text_value(record.get("position_level_from"))
    level_to = text_value(record.get("position_level_to"))

    if holder_status.startswith("1.") and current:
        return current
    if level_from and level_to and level_from != level_to:
        return f"{level_from} - {level_to}"
    return level_from or level_to or current


def build_j18_mapping(ws):
    technical_headers = list(next(ws.iter_rows(
        min_row=J18_TECHNICAL_HEADER_ROW,
        max_row=J18_TECHNICAL_HEADER_ROW,
        values_only=True,
    )))
    header_index = {
        text_value(value): col_no
        for col_no, value in enumerate(technical_headers, start=1)
        if text_value(value)
    }
    missing = sorted(J18_REQUIRED_TECHNICAL_HEADERS - set(header_index))
    if missing:
        raise ValueError(
            "Sheet 2 does not match the verified J18/HROPS schema. "
            f"Missing technical headers on row 6: {', '.join(missing)}"
        )

    mapping = {
        canonical: header_index[technical]
        for canonical, technical in J18_FIELD_MAP.items()
        if technical in header_index
    }
    return mapping, len(technical_headers)


def resolve_or_create_unit(db: Session, record: dict) -> OrganizationalUnit | None:
    unit_id = text_value(record.get("unit_id"))
    unit_name = text_value(record.get("unit_name"))
    province_code = text_value(record.get("province_code"))
    unit_type = text_value(record.get("unit_type")) or "HROPS"
    amphur_name = text_value(record.get("amphur_name"))

    if unit_id:
        unit = db.get(OrganizationalUnit, unit_id)
        if unit:
            if unit_name and unit.unit_name != unit_name:
                unit.unit_name = unit_name
            if unit_type and unit.unit_type_label != unit_type:
                unit.unit_type_label = unit_type
            if amphur_name and unit.amphur_name != amphur_name:
                unit.amphur_name = amphur_name
            return unit

        if province_code and db.get(Province, province_code):
            unit = OrganizationalUnit(
                unit_id=unit_id,
                unit_name=unit_name or unit_id,
                unit_type_label=unit_type,
                province_code=province_code,
                amphur_name=amphur_name,
            )
            db.add(unit)
            db.flush()
            return unit

    if unit_name:
        matches = db.query(OrganizationalUnit).filter(
            OrganizationalUnit.unit_name == unit_name
        ).all()
        if len(matches) == 1:
            return matches[0]

    return None


def process_hrops_workbook(db: Session, run: HropsImportRun) -> dict:
    """
    Verified monthly J18/HROPS import:
    - reads only Sheet 2 ("เขต 1")
    - uses technical headers on row 6
    - data begins on row 7
    - UPSERTs Position Master
    - writes monthly HROPS snapshots
    - never overwrites VacancyCase/VacancyEvent operational workflow data
    """
    run.status = "PROCESSING"
    db.flush()

    summary = {
        "source_sheet": None,
        "technical_header_row": J18_TECHNICAL_HEADER_ROW,
        "columns_detected": 0,
        "rows_seen": 0,
        "inserted": 0,
        "updated": 0,
        "unchanged": 0,
        "occupied": 0,
        "vacant": 0,
        "other_status": 0,
        "skipped": 0,
        "duplicates": 0,
        "errors": [],
    }
    seen_keys = set()

    workbook = load_workbook(run.stored_path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) <= J18_DATA_SHEET_INDEX:
            raise ValueError("J18/HROPS workbook must contain at least 2 sheets.")

        ws = workbook.worksheets[J18_DATA_SHEET_INDEX]
        summary["source_sheet"] = ws.title
        mapping, column_count = build_j18_mapping(ws)
        summary["columns_detected"] = column_count

        for excel_row_no, row in enumerate(
            ws.iter_rows(min_row=J18_TECHNICAL_HEADER_ROW + 1, values_only=True),
            start=J18_TECHNICAL_HEADER_ROW + 1,
        ):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            summary["rows_seen"] += 1

            record = {}
            for canonical, col_no in mapping.items():
                value = row[col_no - 1] if col_no - 1 < len(row) else None
                record[canonical] = json_value(value)

            position_no = text_value(record.get("hrops_position_no"))
            external_position_id = text_value(record.get("external_position_id"))
            position_name = text_value(record.get("position_name_th"))
            holder_status = text_value(record.get("holder_status")) or ""

            if holder_status.startswith("1."):
                summary["occupied"] += 1
            elif holder_status.startswith("2."):
                summary["vacant"] += 1
            else:
                summary["other_status"] += 1

            if not position_no or not position_name:
                summary["skipped"] += 1
                if len(summary["errors"]) < 100:
                    summary["errors"].append({
                        "sheet": ws.title,
                        "row": excel_row_no,
                        "position_no": position_no,
                        "reason": "missing position_code or line_position_name",
                    })
                continue

            source_key = external_position_id or position_no
            if source_key in seen_keys:
                summary["duplicates"] += 1
                continue
            seen_keys.add(source_key)

            unit = resolve_or_create_unit(db, record)
            if not unit:
                summary["skipped"] += 1
                if len(summary["errors"]) < 100:
                    summary["errors"].append({
                        "sheet": ws.title,
                        "row": excel_row_no,
                        "position_no": position_no,
                        "unit_id": text_value(record.get("unit_id")),
                        "unit_name": text_value(record.get("unit_name")),
                        "province_code": text_value(record.get("province_code")),
                        "reason": "position organization could not be matched/created",
                    })
                continue

            stable_chro_id = (
                f"HROPS-{external_position_id}"
                if external_position_id
                else f"HROPS-POS-{position_no}"
            )

            position = db.query(Position).filter(
                Position.chro_position_id == stable_chro_id
            ).first()
            if position is None:
                position = db.query(Position).filter(
                    Position.hrops_position_no == position_no
                ).first()

            position_level = compose_position_level(record)
            created = position is None

            if created:
                chro_id = stable_chro_id
                if db.query(Position).filter(Position.chro_position_id == chro_id).first():
                    chro_id = f"{stable_chro_id}-{uuid4().hex[:8]}"

                position = Position(
                    chro_position_id=chro_id,
                    hrops_position_no=position_no,
                    unit_id=unit.unit_id,
                    position_type=text_value(record.get("position_type")),
                    position_name_th=position_name,
                    position_level=position_level,
                    employment_type=text_value(record.get("employment_type")),
                    is_active=True,
                )
                db.add(position)
                db.flush()
                summary["inserted"] += 1
            else:
                before = (
                    position.hrops_position_no,
                    position.unit_id,
                    position.position_type,
                    position.position_name_th,
                    position.position_level,
                    position.employment_type,
                    position.is_active,
                )
                position.hrops_position_no = position_no
                position.unit_id = unit.unit_id
                position.position_type = text_value(record.get("position_type"))
                position.position_name_th = position_name
                position.position_level = position_level
                position.employment_type = text_value(record.get("employment_type"))
                position.is_active = True
                after = (
                    position.hrops_position_no,
                    position.unit_id,
                    position.position_type,
                    position.position_name_th,
                    position.position_level,
                    position.employment_type,
                    position.is_active,
                )
                if before == after:
                    summary["unchanged"] += 1
                else:
                    summary["updated"] += 1

            # Deliberately store only position/baseline fields in the DB snapshot.
            # Sensitive person-level fields (e.g. citizen_no) remain in the protected source file
            # and are not duplicated into the operational database without a defined use case.
            snapshot_payload = {
                "source_sheet": ws.title,
                "source_row": excel_row_no,
                "external_position_id": external_position_id,
                "hrops_position_no": position_no,
                "unit_id": unit.unit_id,
                "unit_name": unit.unit_name,
                "province_code": text_value(record.get("province_code")),
                "province_name": text_value(record.get("province_name")),
                "position_type": text_value(record.get("position_type")),
                "position_name_th": position_name,
                "position_level_from": text_value(record.get("position_level_from")),
                "position_level_to": text_value(record.get("position_level_to")),
                "current_position_level": text_value(record.get("current_position_level")),
                "position_level": position_level,
                "employment_type": text_value(record.get("employment_type")),
                "holder_status": holder_status,
                "employment_status": text_value(record.get("employment_status")),
                "retirement_date": record.get("retirement_date"),
                "vacancy_date": record.get("vacancy_date"),
                "vacancy_reason": text_value(record.get("vacancy_reason")),
            }

            db.add(HropsPositionSnapshot(
                import_id=run.import_id,
                snapshot_month=run.baseline_month,
                position_uid=position.position_uid,
                hrops_position_no=position_no,
                unit_id=unit.unit_id,
                position_type=text_value(record.get("position_type")),
                position_name_th=position_name,
                position_level=position_level,
                holder_status=holder_status,
                raw_payload=snapshot_payload,
            ))

            if summary["rows_seen"] % 1000 == 0:
                db.flush()

        run.row_count = summary["rows_seen"]
        run.validation_summary = summary
        run.status = "COMPLETED" if summary["skipped"] == 0 else "COMPLETED_WITH_ERRORS"
        db.flush()
        return summary
    finally:
        workbook.close()


app = FastAPI(title="CHRO HR1 Production API", version="3.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "PATCH", "DELETE"],
    allow_headers=["Authorization", "Content-Type"],
)


@app.on_event("startup")
def startup():
    HROPS_STORAGE_DIR.mkdir(parents=True, exist_ok=True)
    Base.metadata.create_all(bind=engine)


@app.get("/healthz")
def healthz():
    return {"status": "ok", "service": "chro-hr1", "version": "3.0.0"}


LOGIN_HTML = """<!doctype html><html lang="th"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>CHRO HR1</title><style>body{font-family:system-ui;margin:0;background:#f5f7fb;color:#172033}.w{max-width:1100px;margin:auto;padding:28px}.c{background:white;border:1px solid #dfe5ef;border-radius:14px;padding:20px;margin:14px 0}input,button{padding:10px;border-radius:8px;border:1px solid #cbd5e1;margin:4px}button{background:#0f5ea8;color:#fff;border:0}.muted{color:#64748b}table{width:100%;border-collapse:collapse}td,th{padding:8px;border-bottom:1px solid #e5e7eb;text-align:left;font-size:14px}.scroll{overflow:auto}</style>
<div class="w"><h1>CHRO HR1</h1><p class="muted">Production Console</p><div id="login" class="c"><form id="f"><input id="u" placeholder="Username" required><input id="p" type="password" placeholder="Password" required><button>Login</button></form><p id="m"></p></div><div id="app" style="display:none"><div class="c"><b id="who"></b> <span id="role"></span> <button id="out">Logout</button><p id="scope" class="muted"></p></div><div id="hropsBox" class="c" style="display:none"><h2>HROPS → Database</h2><p class="muted">เลือกเดือน baseline และไฟล์ .xlsx ระบบจะอ่านไฟล์แล้วอัปเดตฐานข้อมูลทันที</p><input id="hm" type="month"><input id="hf" type="file" accept=".xlsx"><button id="hu">Upload & Update DB</button><pre id="hr"></pre></div><div class="c"><h2>Positions</h2><input id="q" placeholder="Search"><button id="go">Search</button><p id="cnt"></p><div class="scroll"><table><thead><tr><th>CHRO</th><th>HROPS</th><th>Type</th><th>Position</th><th>Level</th><th>Unit</th></tr></thead><tbody id="rows"></tbody></table></div></div></div></div>
<script>const K="chro_token",E=id=>document.getElementById(id),T=()=>sessionStorage.getItem(K);async function A(path,o={}){let h={...(o.headers||{})};if(T())h.Authorization="Bearer "+T();let r=await fetch(path,{...o,headers:h});if(!r.ok){let x=await r.json().catch(()=>({detail:"Request failed"}));throw Error(typeof x.detail==="string"?x.detail:JSON.stringify(x.detail))}return r.json()}async function positions(){let q=encodeURIComponent(E("q").value.trim()),d=await A("/api/positions"+(q?"?q="+q:""));E("cnt").textContent=d.total+" records";E("rows").innerHTML=d.positions.map(p=>`<tr><td>${p.chro_position_id||""}</td><td>${p.hrops_position_no||""}</td><td>${p.position_type||""}</td><td>${p.position_name_th||""}</td><td>${p.position_level||""}</td><td>${p.unit_name||""}</td></tr>`).join("")}async function me(){let x=await A("/api/auth/me");E("who").textContent=x.full_name+" ("+x.username+")";E("role").textContent=x.role;E("scope").textContent=x.unit_id?"Unit: "+x.unit_id:x.province_code?"Province: "+x.province_code:"Region-wide";E("login").style.display="none";E("app").style.display="block";if(["MOPH_ADMIN","REGION_ADMIN"].includes(x.role))E("hropsBox").style.display="block";positions()}E("f").onsubmit=async e=>{e.preventDefault();let b=new URLSearchParams({username:E("u").value,password:E("p").value});try{let r=await fetch("/api/auth/login",{method:"POST",headers:{"Content-Type":"application/x-www-form-urlencoded"},body:b});if(!r.ok)throw Error("Login failed");let x=await r.json();sessionStorage.setItem(K,x.access_token);me()}catch(x){E("m").textContent=x.message}};E("out").onclick=()=>{sessionStorage.removeItem(K);location.reload()};E("go").onclick=positions;E("hu").onclick=async()=>{let f=E("hf").files[0],m=E("hm").value;if(!f||!m){E("hr").textContent="กรุณาเลือกเดือนและไฟล์ .xlsx";return}let b=new FormData();b.append("baseline_month",m+"-01");b.append("file",f);E("hr").textContent="กำลังอ่าน Excel และอัปเดตฐานข้อมูล...";try{let x=await A("/api/hrops/imports",{method:"POST",body:b});E("hr").textContent=JSON.stringify(x.summary,null,2);positions()}catch(x){E("hr").textContent=x.message}};if(T())me().catch(()=>sessionStorage.removeItem(K));</script></html>"""


@app.get("/", response_class=HTMLResponse)
def home():
    return LOGIN_HTML


@app.post("/api/auth/login", response_model=TokenResponse)
def login(request: Request, form: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    user = db.query(User).filter(User.username == form.username.strip()).first()
    if not user or not user.is_active or not pwd_context.verify(form.password, user.password_hash):
        write_audit(db, user if user and user.is_active else None, "AUTH_LOGIN_FAILED", "user", user.user_id if user else None, request)
        db.commit()
        raise HTTPException(status_code=401, detail="Incorrect username or password.")
    user.last_login_at = datetime.now(timezone.utc)
    write_audit(db, user, "AUTH_LOGIN_SUCCESS", "user", user.user_id, request)
    db.commit()
    return TokenResponse(access_token=create_token(user))


@app.get("/api/auth/me")
def me(user: User = Depends(current_user)):
    return {"user_id": user.user_id, "username": user.username, "full_name": user.full_name, "role": user.role, "province_code": user.province_code, "unit_id": user.unit_id}


@app.get("/api/admin/users")
def users(admin: User = Depends(require_roles(UserRole.MOPH_ADMIN, UserRole.REGION_ADMIN)), db: Session = Depends(get_db)):
    rows = db.query(User).order_by(User.role, User.username).all()
    return [{"user_id": x.user_id, "username": x.username, "full_name": x.full_name, "role": x.role, "province_code": x.province_code, "unit_id": x.unit_id, "is_active": x.is_active} for x in rows]


@app.post("/api/admin/users", status_code=201)
def create_user(payload: UserCreate, request: Request, admin: User = Depends(require_roles(UserRole.MOPH_ADMIN, UserRole.REGION_ADMIN)), db: Session = Depends(get_db)):
    if db.query(User).filter(User.username == payload.username.strip()).first():
        raise HTTPException(status_code=409, detail="Username already exists.")
    if admin.role == UserRole.REGION_ADMIN and payload.role == UserRole.MOPH_ADMIN:
        raise HTTPException(status_code=403, detail="REGION_ADMIN cannot create MOPH_ADMIN.")
    if payload.role == UserRole.PROVINCE_ADMIN and not payload.province_code:
        raise HTTPException(status_code=422, detail="province_code required.")
    if payload.role == UserRole.HOSPITAL_HR and not payload.unit_id:
        raise HTTPException(status_code=422, detail="unit_id required.")
    if payload.province_code and not db.get(Province, payload.province_code):
        raise HTTPException(status_code=422, detail="Unknown province_code.")
    if payload.unit_id and not db.get(OrganizationalUnit, payload.unit_id):
        raise HTTPException(status_code=422, detail="Unknown unit_id.")
    user = User(username=payload.username.strip(), password_hash=hash_password(payload.password), full_name=payload.full_name.strip(), role=payload.role, province_code=payload.province_code, unit_id=payload.unit_id)
    db.add(user); db.flush()
    write_audit(db, admin, "USER_CREATED", "user", user.user_id, request, after={"username": user.username, "role": user.role.value, "province_code": user.province_code, "unit_id": user.unit_id})
    db.commit()
    return {"user_id": user.user_id, "username": user.username, "role": user.role}


@app.post("/api/positions", status_code=201)
def create_position(payload: PositionCreate, request: Request, user: User = Depends(require_roles(UserRole.MOPH_ADMIN, UserRole.REGION_ADMIN)), db: Session = Depends(get_db)):
    if db.query(Position).filter(Position.chro_position_id == payload.chro_position_id).first():
        raise HTTPException(status_code=409, detail="chro_position_id already exists.")
    if not db.get(OrganizationalUnit, payload.unit_id):
        raise HTTPException(status_code=422, detail="Unknown unit_id.")
    p = Position(**payload.model_dump())
    db.add(p); db.flush()
    write_audit(db, user, "POSITION_CREATED", "position", p.position_uid, request, after=payload.model_dump())
    db.commit()
    return {"position_uid": p.position_uid, "chro_position_id": p.chro_position_id}


@app.get("/api/positions")
def positions(q: str | None = None, user: User = Depends(current_user), db: Session = Depends(get_db)):
    query = db.query(Position, OrganizationalUnit).join(OrganizationalUnit, Position.unit_id == OrganizationalUnit.unit_id)
    if user.role == UserRole.PROVINCE_ADMIN:
        query = query.filter(OrganizationalUnit.province_code == user.province_code)
    elif user.role == UserRole.HOSPITAL_HR:
        query = query.filter(Position.unit_id == user.unit_id)
    elif user.role not in READ_ALL:
        query = query.filter(Position.position_uid == "__denied__")
    if q:
        term = f"%{q.strip()}%"
        query = query.filter(or_(Position.chro_position_id.ilike(term), Position.hrops_position_no.ilike(term), Position.position_name_th.ilike(term), Position.position_level.ilike(term), OrganizationalUnit.unit_name.ilike(term)))
    rows = query.order_by(OrganizationalUnit.province_code, OrganizationalUnit.unit_name).limit(2000).all()
    return {"total": len(rows), "positions": [{"position_uid": p.position_uid, "chro_position_id": p.chro_position_id, "hrops_position_no": p.hrops_position_no, "position_type": p.position_type, "position_name_th": p.position_name_th, "position_level": p.position_level, "employment_type": p.employment_type, "unit_id": u.unit_id, "unit_name": u.unit_name, "province_code": u.province_code} for p, u in rows]}


@app.post("/api/vacancies", status_code=201)
def create_vacancy(payload: VacancyCreate, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    p = db.get(Position, payload.position_uid)
    if not p:
        raise HTTPException(status_code=404, detail="Position not found.")
    assert_unit_scope(db, user, payload.responsible_unit_id, write=True)
    assert_unit_scope(db, user, p.unit_id, write=True)
    if db.query(VacancyCase).filter(VacancyCase.case_no == payload.case_no).first():
        raise HTTPException(status_code=409, detail="case_no exists.")
    c = VacancyCase(**payload.model_dump(), created_by_user_id=user.user_id)
    db.add(c); db.flush()
    db.add(VacancyEvent(case_id=c.case_id, event_type="CASE_CREATED", to_status=c.status.value, milestone=c.current_milestone, notes=c.remarks, actor_user_id=user.user_id))
    write_audit(db, user, "VACANCY_CREATED", "vacancy_case", c.case_id, request, after={"case_no": c.case_no, "status": c.status.value, "milestone": c.current_milestone})
    db.commit()
    return {"case_id": c.case_id, "case_no": c.case_no, "status": c.status}


@app.get("/api/vacancies")
def vacancies(user: User = Depends(current_user), db: Session = Depends(get_db)):
    query = db.query(VacancyCase, Position, OrganizationalUnit).join(Position, VacancyCase.position_uid == Position.position_uid).join(OrganizationalUnit, Position.unit_id == OrganizationalUnit.unit_id)
    if user.role == UserRole.PROVINCE_ADMIN:
        query = query.filter(OrganizationalUnit.province_code == user.province_code)
    elif user.role == UserRole.HOSPITAL_HR:
        query = query.filter(VacancyCase.responsible_unit_id == user.unit_id)
    elif user.role not in READ_ALL:
        query = query.filter(VacancyCase.case_id == "__denied__")
    rows = query.order_by(VacancyCase.updated_at.desc()).limit(2000).all()
    return {"total": len(rows), "vacancies": [{"case_id": c.case_id, "case_no": c.case_no, "chro_position_id": p.chro_position_id, "position_name_th": p.position_name_th, "position_level": p.position_level, "unit_name": u.unit_name, "province_code": u.province_code, "vacant_date": c.vacant_date, "vacant_reason": c.vacant_reason, "status": c.status, "current_milestone": c.current_milestone, "retirement_use_approved": c.retirement_use_approved, "retirement_approval_doc_no": c.retirement_approval_doc_no, "retirement_use_from_date": c.retirement_use_from_date, "remarks": c.remarks} for c, p, u in rows]}


@app.post("/api/vacancies/{case_id}/events", status_code=201)
def vacancy_event(case_id: str, payload: VacancyEventCreate, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    c = db.get(VacancyCase, case_id)
    if not c:
        raise HTTPException(status_code=404, detail="Vacancy case not found.")
    assert_unit_scope(db, user, c.responsible_unit_id, write=True)
    before = {"status": c.status.value, "milestone": c.current_milestone, "retirement_use_approved": c.retirement_use_approved}
    old_status = c.status.value
    if payload.to_status is not None:
        c.status = payload.to_status
        if c.status == VacancyStatus.FILLED:
            c.closed_at = datetime.now(timezone.utc)
    if payload.milestone is not None:
        c.current_milestone = payload.milestone
    if payload.retirement_use_approved is not None:
        c.retirement_use_approved = payload.retirement_use_approved
    if payload.retirement_approval_doc_no is not None:
        c.retirement_approval_doc_no = payload.retirement_approval_doc_no
    if payload.retirement_approval_doc_date is not None:
        c.retirement_approval_doc_date = payload.retirement_approval_doc_date
    if payload.retirement_use_from_date is not None:
        c.retirement_use_from_date = payload.retirement_use_from_date
    e = VacancyEvent(case_id=c.case_id, event_type=payload.event_type, from_status=old_status, to_status=c.status.value, milestone=c.current_milestone, reference_doc_no=payload.reference_doc_no, notes=payload.notes, actor_user_id=user.user_id)
    db.add(e); db.flush()
    write_audit(db, user, "VACANCY_EVENT_ADDED", "vacancy_case", c.case_id, request, before=before, after={"event_type": payload.event_type, "status": c.status.value, "milestone": c.current_milestone, "retirement_use_approved": c.retirement_use_approved, "retirement_approval_doc_no": c.retirement_approval_doc_no, "retirement_use_from_date": str(c.retirement_use_from_date) if c.retirement_use_from_date else None})
    db.commit()
    return {"event_id": e.event_id, "status": c.status, "milestone": c.current_milestone}


@app.get("/api/hrops/imports")
def hrops_imports(user: User = Depends(current_user), db: Session = Depends(get_db)):
    if user.role not in READ_ALL:
        raise HTTPException(status_code=403, detail="HROPS baseline history not available to this role.")
    rows = db.query(HropsImportRun).order_by(HropsImportRun.baseline_month.desc(), HropsImportRun.created_at.desc()).limit(120).all()
    return [{"import_id": x.import_id, "baseline_month": x.baseline_month, "original_filename": x.original_filename, "sha256": x.sha256, "file_size_bytes": x.file_size_bytes, "status": x.status, "row_count": x.row_count, "validation_summary": x.validation_summary, "created_at": x.created_at} for x in rows]


@app.post("/api/hrops/imports", status_code=201)
async def upload_hrops(request: Request, baseline_month: date = Form(...), notes: str | None = Form(None), file: UploadFile = File(...), user: User = Depends(require_roles(UserRole.MOPH_ADMIN, UserRole.REGION_ADMIN)), db: Session = Depends(get_db)):
    filename = file.filename or "hrops.xlsx"
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="Only .xlsx accepted.")
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(filename).name)
    folder = HROPS_STORAGE_DIR / baseline_month.strftime("%Y-%m")
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"{uuid4()}_{safe}"
    limit = MAX_HROPS_UPLOAD_MB * 1024 * 1024
    sha = hashlib.sha256(); total = 0
    try:
        with target.open("wb") as out:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > limit:
                    raise HTTPException(status_code=413, detail=f"File exceeds {MAX_HROPS_UPLOAD_MB} MB.")
                sha.update(chunk); out.write(chunk)
    except Exception:
        target.unlink(missing_ok=True)
        raise

    digest = sha.hexdigest()
    duplicate = db.query(HropsImportRun).filter(HropsImportRun.baseline_month == baseline_month, HropsImportRun.sha256 == digest).first()
    if duplicate:
        target.unlink(missing_ok=True)
        raise HTTPException(status_code=409, detail={"message": "Duplicate HROPS file for month.", "import_id": duplicate.import_id})

    run = HropsImportRun(
        baseline_month=baseline_month,
        original_filename=filename,
        stored_path=str(target),
        sha256=digest,
        file_size_bytes=total,
        status="RECEIVED",
        notes=notes,
        uploaded_by_user_id=user.user_id,
    )
    db.add(run)
    db.flush()
    write_audit(
        db, user, "HROPS_FILE_RECEIVED", "hrops_import", run.import_id, request,
        after={"baseline_month": str(baseline_month), "filename": filename, "sha256": digest, "size": total},
    )
    db.commit()
    run = db.get(HropsImportRun, run.import_id)

    try:
        summary = process_hrops_workbook(db, run)
        write_audit(
            db, user, "HROPS_DB_UPDATED", "hrops_import", run.import_id, request,
            after={
                "baseline_month": str(baseline_month),
                "filename": filename,
                "sha256": digest,
                "size": total,
                **{k: v for k, v in summary.items() if k != "errors"},
            },
        )
        db.commit()
    except Exception as exc:
        db.rollback()
        failed = db.get(HropsImportRun, run.import_id)
        if failed:
            failed.status = "FAILED"
            failed.validation_summary = {"error": str(exc)}
            db.commit()
        raise HTTPException(status_code=422, detail=f"HROPS import failed: {exc}")

    return {
        "import_id": run.import_id,
        "status": run.status,
        "sha256": run.sha256,
        "file_size_bytes": run.file_size_bytes,
        "summary": summary,
        "database_updated": True,
    }


@app.get("/api/audit")
def audit_logs(limit: int = 200, user: User = Depends(require_roles(UserRole.MOPH_ADMIN, UserRole.REGION_ADMIN, UserRole.AUDITOR)), db: Session = Depends(get_db)):
    rows = db.query(AuditLog).order_by(AuditLog.created_at.desc()).limit(max(1, min(limit, 1000))).all()
    return [{"audit_id": x.audit_id, "actor_user_id": x.actor_user_id, "action": x.action, "entity_type": x.entity_type, "entity_id": x.entity_id, "request_path": x.request_path, "ip_address": x.ip_address, "before_json": x.before_json, "after_json": x.after_json, "created_at": x.created_at} for x in rows]


PROVINCES = [("50","เชียงใหม่"),("51","ลำพูน"),("52","ลำปาง"),("54","แพร่"),("55","น่าน"),("56","พะเยา"),("57","เชียงราย"),("58","แม่ฮ่องสอน")]
UNITS = [
("U5001","รพศ.นครพิงค์","รพศ.","50","แม่ริม"),("U5002","สสจ.เชียงใหม่","สสจ.","50","เมืองเชียงใหม่"),
("U5101","รพท.ลำพูน","รพท.","51","เมืองลำพูน"),("U5102","สสจ.ลำพูน","สสจ.","51","เมืองลำพูน"),
("U5201","รพศ.ลำปาง","รพศ.","52","เมืองลำปาง"),("U5202","สสจ.ลำปาง","สสจ.","52","เมืองลำปาง"),
("U5401","รพท.แพร่","รพท.","54","เมืองแพร่"),("U5402","สสจ.แพร่","สสจ.","54","เมืองแพร่"),
("U5501","รพท.น่าน","รพท.","55","เมืองน่าน"),("U5502","สสจ.น่าน","สสจ.","55","เมืองน่าน"),
("U5601","รพท.พะเยา","รพท.","56","เมืองพะเยา"),("U5602","รพท.เชียงคำ","รพท.","56","เชียงคำ"),("U5603","สสจ.พะเยา","สสจ.","56","เมืองพะเยา"),
("U5701","รพศ.เชียงรายประชานุเคราะห์","รพศ.","57","เมืองเชียงราย"),("U5702","สสจ.เชียงราย","สสจ.","57","เมืองเชียงราย"),
("U5801","รพท.ศรีสังวาลย์","รพท.","58","เมืองแม่ฮ่องสอน"),("U5802","สสจ.แม่ฮ่องสอน","สสจ.","58","เมืองแม่ฮ่องสอน"),
]


def init_reference():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        for code, name in PROVINCES:
            if not db.get(Province, code):
                db.add(Province(province_code=code, province_name_th=name, health_region="1"))
        db.flush()
        for unit_id, name, typ, province, amphur in UNITS:
            if not db.get(OrganizationalUnit, unit_id):
                db.add(OrganizationalUnit(unit_id=unit_id, unit_name=name, unit_type_label=typ, province_code=province, amphur_name=amphur))
        db.commit()
    finally:
        db.close()


def bootstrap_admin(username: str, password: str, full_name: str, role: UserRole):
    if role not in {UserRole.MOPH_ADMIN, UserRole.REGION_ADMIN}:
        raise SystemExit("Bootstrap role must be MOPH_ADMIN or REGION_ADMIN.")
    init_reference()
    db = SessionLocal()
    try:
        if db.query(User).filter(User.username == username).first():
            raise SystemExit("Username already exists.")
        db.add(User(username=username, password_hash=hash_password(password), full_name=full_name, role=role, is_active=True))
        db.commit()
    finally:
        db.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("init-reference")
    a = sub.add_parser("bootstrap-admin")
    a.add_argument("--username", required=True)
    a.add_argument("--password", required=True)
    a.add_argument("--full-name", required=True)
    a.add_argument("--role", choices=[UserRole.MOPH_ADMIN.value, UserRole.REGION_ADMIN.value], default=UserRole.REGION_ADMIN.value)
    args = parser.parse_args()
    if args.command == "init-reference":
        init_reference()
        print("Region 1 reference data ready.")
    else:
        bootstrap_admin(args.username, args.password, args.full_name, UserRole(args.role))
        print(f"Created {args.role}: {args.username}")
